"""Word报告生成服务"""
import io
from datetime import datetime
from typing import Dict, Any, Optional, List
from loguru import logger

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    Document = None
    Pt = None
    RGBColor = None
    WD_ALIGN_PARAGRAPH = None
    Inches = None
    logger.warning("python-docx 未安装，Word报告功能不可用")


class WordReportService:
    """Word报告生成服务"""

    def __init__(self):
        self.available = DOCX_AVAILABLE

    def create_report(
        self,
        results: Dict[str, Any],
        current_analysis: Optional[Dict[str, Any]] = None,
        include_charts: bool = False,
        progress_callback=None,
    ) -> Optional[io.BytesIO]:
        """生成综合分析报告"""
        if not self.available:
            logger.error("python-docx 不可用")
            return None

        def update_progress(msg: str):
            if progress_callback:
                progress_callback(msg)

        try:
            update_progress("初始化文档...")
            doc = Document()

            # 标题
            title = doc.add_heading("商品期货 AI 分析报告", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 报告信息
            date_str = current_analysis.get("analysis_date", datetime.now().strftime("%Y-%m-%d")) if current_analysis else datetime.now().strftime("%Y-%m-%d")
            commodities = list(results.keys())
            info = doc.add_paragraph()
            info.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = info.add_run(f"分析日期: {date_str} | 品种: {', '.join(commodities)}")
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(128, 128, 128)

            doc.add_paragraph()

            for i, (commodity, result) in enumerate(results.items()):
                update_progress(f"生成 {commodity} 报告 ({i+1}/{len(results)})...")
                self._add_commodity_section(doc, commodity, result, include_charts)

            update_progress("最终整理...")

            # 保存到内存
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            update_progress("完成!")
            return buffer

        except Exception as e:
            logger.error(f"生成Word报告失败: {e}")
            return None

    def _add_commodity_section(self, doc: Document, commodity: str, result: Dict[str, Any], include_charts: bool):
        """添加单个品种的分析章节"""
        doc.add_heading(f"{commodity} 分析报告", level=1)

        # 基本信息
        analysis_date = result.get("analysis_date", "未知")
        doc.add_paragraph(f"分析日期: {analysis_date}")

        # 各模块分析结果
        modules = result.get("modules", {})
        if modules:
            doc.add_heading("模块分析结果", level=2)

            module_names = {
                "inventory": "库存仓单分析",
                "positioning": "持仓席位分析",
                "term_structure": "期限结构分析",
                "technical": "技术面分析",
                "basis": "基差分析",
                "news": "新闻分析",
            }

            for module_key, module_result in modules.items():
                display_name = module_names.get(module_key, module_key)
                doc.add_heading(display_name, level=3)

                status = module_result.get("status", "unknown")
                confidence = module_result.get("confidence_score", 0)
                doc.add_paragraph(f"状态: {status} | 信心度: {confidence:.0%}")

                summary = module_result.get("analysis_summary", "")
                if summary:
                    doc.add_paragraph(summary)

                findings = module_result.get("key_findings", [])
                if findings:
                    for finding in findings:
                        doc.add_paragraph(finding, style="List Bullet")

        # 辩论结果
        debate = result.get("debate")
        if debate:
            doc.add_heading("多空辩论", level=2)
            doc.add_paragraph(f"辩论轮数: {debate.get('rounds', 0)}")
            doc.add_paragraph(f"多头评分: {debate.get('bull_score', 0):.0%}")
            doc.add_paragraph(f"空头评分: {debate.get('bear_score', 0):.0%}")
            doc.add_paragraph(f"胜方: {debate.get('winner', '未知')}")

        # 交易员建议
        trader = result.get("trader")
        if trader:
            doc.add_heading("交易员建议", level=2)
            doc.add_paragraph(f"方向: {trader.get('direction', '未知')}")
            doc.add_paragraph(f"信心度: {trader.get('confidence', 0):.0%}")
            doc.add_paragraph(f"仓位建议: {trader.get('position_size', '未知')}")

        # 风控意见
        risk = result.get("risk_management")
        if risk:
            doc.add_heading("风控意见", level=2)
            doc.add_paragraph(f"是否批准: {'是' if risk.get('approved') else '否'}")
            doc.add_paragraph(f"风险等级: {risk.get('risk_level', '未知')}")
            doc.add_paragraph(f"最大仓位: {risk.get('max_position', 0):.0%}")

        # 最终决策
        decision = result.get("executive_decision", {})
        if decision:
            doc.add_heading("最终决策", level=2)
            doc.add_paragraph(f"决策: {decision.get('final_decision', '未知')}")
            doc.add_paragraph(f"信心等级: {decision.get('confidence_level', '未知')}")
            doc.add_paragraph(f"方向判断: {decision.get('directional_view', '未知')}")
            doc.add_paragraph(f"方向信心: {decision.get('directional_confidence', 0):.0%}")

            reasoning = decision.get("reasoning", "")
            if reasoning:
                doc.add_paragraph(reasoning)

            action_items = decision.get("action_items", [])
            if action_items:
                doc.add_heading("操作要点", level=3)
                for item in action_items:
                    doc.add_paragraph(item, style="List Bullet")

        doc.add_page_break()


# 全局单例
word_report_service = WordReportService()
